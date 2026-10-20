#!/usr/bin/env python3
"""Build the complete Gravity AI Prompt workbook from the frozen 250-prompt
baseline plus non-destructively appended, evidence-controlled additions.

Outputs are a searchable JSON index, matching reader text, editable Word book
and searchable printable PDF. Never regenerate or modify the historical
250-prompt snapshot or the older 215-prompt files.
"""
from __future__ import annotations
from collections import Counter
from datetime import date
from html import escape
from pathlib import Path
import json
import re

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import KeepTogether, PageBreak, Paragraph, SimpleDocTemplate, Spacer
from pypdf import PdfReader

ROOT = Path(__file__).resolve().parents[1]
INDEX = ROOT / 'site/assets/data'
READER = ROOT / 'site/assets/reader/ai-prompts.txt'
OUT = ROOT / 'site/assets/free-resources/ai'
COUNT = 325
BASE = 250
STEM = 'RESTAURANT_AI_PROMPTS_AND_SYSTEMS_WORKBOOK_325_PROMPTS_2026-10-20'
NAVY = colors.HexColor('#173340')
GOLD = colors.HexColor('#AD8035')
TEXT = colors.HexColor('#243745')
MUTED = colors.HexColor('#57646A')
PAPER = colors.HexColor('#F7F2E9')

def fetch_rows():
    baseline = json.loads((INDEX/'restaurant-ai-prompts-250.json').read_text(encoding='utf-8'))
    additions = []
    for first, last in ((251,275),(276,300),(301,325)):
        path = ROOT/f'scripts/ai-prompts-{first}-{last}.json'
        rows = json.loads(path.read_text(encoding='utf-8'))
        assert len(rows) == last-first+1, f'Wrong count in {path}'
        additions.extend(rows)
    assert len(baseline) == BASE, f'Frozen baseline changed: {len(baseline)}'
    rows = baseline + additions
    assert len(rows) == COUNT, f'Expected {COUNT}, saw {len(rows)}'
    ids = [r['id'] for r in rows]
    names = [r['title'].strip().casefold() for r in rows]
    assert len(set(ids)) == COUNT, 'Duplicate IDs'
    assert len(set(names)) == COUNT, 'Duplicate titles'
    for i, r in enumerate(rows, 1):
        assert r['number'] == i and r['id'] == f'GRV-AI-{i:03d}', f'Broken numbering: {r}'
        for key in ('title','prompt','manager_check'):
            assert isinstance(r.get(key), str) and len(r[key].strip()) > 8, f'Missing {key} at {i}'
        assert isinstance(r.get('category'), str) and len(r['category'].strip()) > 2, f'Missing category at {i}'
        assert len(r['prompt']) >= 100, f'Prompt {i} too short for a useful copy/paste'
    assert rows[249] == baseline[-1] and rows[250]['id'] == 'GRV-AI-251', 'Baseline not preserved'
    return rows

def group_sections(rows):
    ordered=[]
    groups={}
    for r in rows:
        cat=r['category']
        if cat not in groups: groups[cat]=[];ordered.append(cat)
        groups[cat].append(r)
    return ordered, groups

def reader_text(rows, ordered, groups):
    top=[
        'GRAVITY','THE RESTAURANT LEADERSHIP SYSTEM','',
        'Restaurant AI Prompts &','Systems Workbook','',
        f'{len(rows)} practical copy-and-paste prompts for restaurant leadership,',
        'service, Work Study Lab, SOP Foundry, Shift Lab, Inventory Center,',
        'Testing Room, Windows software, governance, and verified improvement.','',
        'THE GRAVITY SYSTEM AI COMPANION','',
        'Use AI to organize the work. Keep management judgment, real evidence,',
        'approved procedures, customer privacy and verified results in control.','',
        'THOMAS BUTLER | THOMAS MONROE BOOKS',
        'Copyright 2026 Thomas Butler. All rights reserved.',
        'These are reusable draft prompts, not legal, clinical, employment',
        'or food-safety determinations. Managers verify and approve decisions.','',
        'TABLE OF CONTENTS — CATEGORIES',''
    ]
    for cat in ordered:
        r=groups[cat]
        top.append(f'{cat} — {len(r)} prompts (#{r[0]["number"]}–#{r[-1]["number"]})'
                   if all(x['number'] == r[0]['number']+i for i,x in enumerate(r))
                   else f'{cat} — {len(r)} prompts')
    top+=['','HOW TO USE THE WORKBOOK','',
          '1. Pick the prompt that matches one specific problem.',
          '2. Replace the bracketed fields with verified restaurant context.',
          '3. Ask for missing evidence and review the manager check.',
          '4. Approve, measure and verify any operational decision separately.',
          '5. Protect employee, customer, recipe and credential data.','',
          'VERSION CONTROL','',
          'Prompts 001–250 retain the existing controlled digital text.',
          'Prompts 251–325 are appended additions based on later Gravity work.',
          'Earlier editions remain archived, not silently replaced.','']
    for cat in ordered:
        top.extend(['', '='*66, cat, '='*66])
        for r in groups[cat]:
            top.extend(['',f'PROMPT {r["number"]:03d} • {r["title"]}',f'ID: {r["id"]}','',r['prompt'],'',
                        f'Manager check: {r["manager_check"]}'])
    return '\n'.join(top)+'\n'

def set_docx_cell_shading(par, color='F4EEE3'):
    pPr=par._p.get_or_add_pPr()
    shd=OxmlElement('w:shd')
    shd.set(qn('w:fill'),color)
    pPr.append(shd)

def set_page_field(par):
    par.add_run('PAGE ')
    field=OxmlElement('w:fldSimple')
    field.set(qn('w:instr'),'PAGE')
    par._p.append(field)

def build_docx(rows,ordered,groups,path):
    doc=Document()
    sec=doc.sections[0]
    sec.top_margin=Inches(.78);sec.bottom_margin=Inches(.68)
    sec.left_margin=Inches(.80);sec.right_margin=Inches(.80)
    sec.header_distance=Inches(.38);sec.footer_distance=Inches(.37)
    normal=doc.styles['Normal']
    normal.font.name='Aptos';normal.font.size=Pt(10)
    normal.font.color.rgb=RGBColor(36,55,69)
    normal.paragraph_format.space_after=Pt(8)
    for label, size, before, after, color in [
        ('Title',27,0,15,RGBColor(23,51,64)),
        ('Heading 1',15,17,10,RGBColor(23,51,64)),
        ('Heading 2',11,11,5,RGBColor(23,51,64)),
    ]:
        st=doc.styles[label]
        st.font.name='Aptos Display' if label=='Title' else 'Aptos'
        st.font.size=Pt(size)
        st.font.bold=True
        st.font.color.rgb=color
        st.paragraph_format.space_before=Pt(before)
        st.paragraph_format.space_after=Pt(after)
        st.paragraph_format.keep_with_next=True
    if 'Manager Check' not in doc.styles:
        st=doc.styles.add_style('Manager Check',WD_STYLE_TYPE.PARAGRAPH)
        st.base_style=doc.styles['Normal']
    check=doc.styles['Manager Check']
    check.font.size=Pt(9)
    check.font.italic=True
    check.paragraph_format.left_indent=Inches(.14)
    check.paragraph_format.right_indent=Inches(.08)
    check.paragraph_format.space_before=Pt(6)
    check.paragraph_format.space_after=Pt(14)
    header=sec.header.paragraphs[0]
    header.text='THE GRAVITY SYSTEM   /   THOMAS MONROE BOOKS'
    header.style=doc.styles['Caption']
    footer=sec.footer.paragraphs[0]
    footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run('THOMAS BUTLER  ·  RESTAURANT AI WORKBOOK  ·  ')
    set_page_field(footer)
    doc.core_properties.title='Restaurant AI Prompts & Systems Workbook — 325 Prompts'
    doc.core_properties.author='Thomas Butler'
    doc.core_properties.subject='The Gravity Restaurant Leadership System'
    doc.add_paragraph('THE GRAVITY RESTAURANT LEADERSHIP SYSTEM','Subtitle')
    doc.add_heading('Restaurant AI Prompts & Systems Workbook',0)
    h=doc.add_paragraph(f'{len(rows)} practical prompts — complete digital edition')
    h.style='Subtitle'
    doc.add_paragraph('Restaurant operations · Work Study Lab · SOP control · Staffing · Inventory · Testing & Windows · Owner–GM governance')
    doc.add_paragraph('THOMAS BUTLER  |  THOMAS MONROE BOOKS')
    doc.add_paragraph('October 2026 digital update  •  Prompts 001–250 preserved; 251–325 appended.')
    doc.add_paragraph('These prompts support professional judgment and working drafts. Managers must verify source records, employee protections, safety standards and operational approval before implementation.')
    doc.add_page_break()
    doc.add_heading('How to use this workbook',level=1)
    for step in [
        'Choose one prompt for a specific problem or decision.',
        'Replace [bracketed fields] with verified facts; omit secrets and unnecessary personal data.',
        'Review the Manager check and request missing evidence.',
        'Keep drafts, approvals, measurements and follow-up separate.',
        'Use the online search to locate prompt IDs and newer category additions.'
    ]:doc.add_paragraph(step,style='List Number')
    doc.add_heading('Category guide',level=1)
    for cat in ordered:doc.add_paragraph(f'{cat.title()} — {len(groups[cat])} prompts',style='List Bullet')
    doc.add_heading('Version control',level=1)
    doc.add_paragraph('The existing 250-prompt digital index is the preserved baseline. Seventy-five newly developed prompts were added without deleting or renumbering a prior prompt. Older downloadable editions remain archival and are not labeled as the current edition.')
    doc.add_page_break()
    for cat in ordered:
        doc.add_heading(cat.title(),level=1)
        for r in groups[cat]:
            h=doc.add_heading(f'{r["number"]:03d} · {r["title"].title()}',level=2)
            h.paragraph_format.keep_with_next=True
            ident=doc.add_paragraph(r['id'])
            ident.style=doc.styles['Caption']
            ident.paragraph_format.space_after=Pt(5)
            doc.add_paragraph(r['prompt'])
            p=doc.add_paragraph(style='Manager Check')
            set_docx_cell_shading(p)
            x=p.add_run('MANAGER CHECK  ')
            x.bold=True
            p.add_run(r['manager_check'])
    OUT.mkdir(parents=True,exist_ok=True)
    doc.save(str(path))
    assert path.stat().st_size>50000, 'DOCX unexpectedly empty'

def register_font():
    font=Path('/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf')
    bold=Path('/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf')
    assert font.exists() and bold.exists(), 'Missing PDF Unicode fonts'
    pdfmetrics.registerFont(TTFont('GravitySans',str(font)))
    pdfmetrics.registerFont(TTFont('GravitySans-Bold',str(bold)))

def cover_footer(canvas, doc):
    canvas.saveState()
    canvas.setStrokeColor(GOLD)
    canvas.setLineWidth(.8)
    canvas.line(.70*inch,.62*inch,7.80*inch,.62*inch)
    canvas.setFont('GravitySans',7.5)
    canvas.setFillColor(MUTED)
    canvas.drawString(.72*inch,.43*inch,'THOMAS MONROE BOOKS  /  THE GRAVITY SYSTEM  /  OCTOBER 2026')
    canvas.drawRightString(7.77*inch,.43*inch,str(doc.page))
    canvas.restoreState()

def pdf_markup(s):
    return escape(str(s)).replace('\n\n','<br/><br/>').replace('\n','<br/>')

def build_pdf(rows,ordered,groups,path):
    register_font()
    ss=getSampleStyleSheet()
    title=ParagraphStyle('GravityTitle',fontName='GravitySans-Bold',fontSize=22,leading=30,textColor=NAVY,spaceAfter=17)
    cover=ParagraphStyle('GravityCover',fontName='GravitySans',fontSize=12,leading=19,textColor=TEXT,spaceAfter=14)
    body=ParagraphStyle('GravityBody',fontName='GravitySans',fontSize=8.9,leading=14.3,textColor=TEXT,spaceAfter=11)
    cat=ParagraphStyle('GravityCat',fontName='GravitySans-Bold',fontSize=13,leading=20,textColor=NAVY,spaceBefore=16,spaceAfter=13,keepWithNext=True)
    head=ParagraphStyle('GravityPrompt',fontName='GravitySans-Bold',fontSize=10,leading=15,textColor=NAVY,spaceBefore=11,spaceAfter=5,keepWithNext=True)
    hint=ParagraphStyle('GravityHint',fontName='GravitySans',fontSize=8.5,leading=13,textColor=MUTED,spaceBefore=5,spaceAfter=14,leftIndent=10,borderColor=GOLD,borderPadding=5,borderWidth=.7)
    small=ParagraphStyle('GravitySmall',fontName='GravitySans',fontSize=8.4,leading=13,textColor=TEXT,spaceAfter=4)
    subtitle=ParagraphStyle('GravitySubtitle',fontName='GravitySans-Bold',fontSize=9.4,leading=14,textColor=GOLD,spaceAfter=15)
    story=[]
    story.append(Spacer(1,.85*inch))
    story.append(Paragraph('THE GRAVITY RESTAURANT LEADERSHIP SYSTEM',subtitle))
    story.append(Paragraph('Restaurant AI Prompts<br/>&amp; Systems Workbook',title))
    story.append(Paragraph(f'{len(rows)} practical copy-and-paste restaurant leadership prompts',cover))
    story.append(Paragraph('Work Study Lab  ·  SOP Foundry  ·  Shift Lab  ·  Inventory Center<br/>Testing &amp; Windows  ·  Owner–GM governance  ·  Verified improvement',cover))
    story.append(Spacer(1,.42*inch))
    story.append(Paragraph('THOMAS BUTLER  |  THOMAS MONROE BOOKS',subtitle))
    story.append(Paragraph('October 2026 digital update · Prompts 001–250 preserved; 251–325 appended.',cover))
    story.append(Paragraph('The prompts organize working drafts and questions; they do not replace owner, manager, safety, HR or legal judgment. Review all applicable standards and verify evidence before acting.',body))
    story.append(PageBreak())
    story.append(Paragraph('How to use this workbook',cat))
    for n,step in enumerate([
        'Choose one prompt for the actual problem or decision.',
        'Replace bracketed fields with verified restaurant facts; protect confidential information.',
        'Read the Manager check, ask about missing evidence, and record assumptions.',
        'Keep drafts separate from management approvals and observed verification.',
        'Search by prompt ID or category at thomasmonroebooks.com/restaurant-ai-prompt-search.',
    ],1):story.append(Paragraph(f'{n}. {pdf_markup(step)}',body))
    story.append(Paragraph('Contents — categories',cat))
    for c in ordered:story.append(Paragraph(f'{pdf_markup(c.title())}  ·  {len(groups[c])} prompts',small))
    story.append(Spacer(1,.15*inch))
    story.append(Paragraph('Version control',cat))
    story.append(Paragraph('The existing controlled 250-prompt index has been kept intact. Seventy-five new prompts are appended. Older downloadable files remain archived; this PDF and matching Word edition contain the complete current set.',body))
    story.append(PageBreak())
    for category in ordered:
        story.append(Paragraph(pdf_markup(category.title()),cat))
        for r in groups[category]:
            h=Paragraph(f'{r["number"]:03d}  ·  {pdf_markup(r["title"].title())}',head)
            ref=Paragraph(pdf_markup(r['id']),small)
            prompt=Paragraph(pdf_markup(r['prompt']),body)
            manager=Paragraph('<b>MANAGER CHECK</b>  '+pdf_markup(r['manager_check']),hint)
            story.extend([KeepTogether([h,ref]),prompt,manager])
    doc=SimpleDocTemplate(str(path),pagesize=letter,leftMargin=.75*inch,rightMargin=.75*inch,
                          topMargin=.74*inch,bottomMargin=.78*inch,
                          title='Restaurant AI Prompts & Systems Workbook — 325 Prompts',
                          author='Thomas Butler',subject='The Gravity Restaurant Leadership System')
    doc.build(story,onFirstPage=cover_footer,onLaterPages=cover_footer)
    pages=PdfReader(str(path))
    text='\n'.join(p.extract_text() or '' for p in pages.pages)
    for i in (1,200,250,251,275,300,325):
        assert f'GRV-AI-{i:03d}' in text, f'Prompt {i} missing from PDF'
    assert path.stat().st_size>100000, 'PDF unexpectedly small'
    print('PDF_PAGES='+str(len(pages.pages)))

def verify_docx(path):
    d=Document(str(path))
    text='\n'.join(p.text for p in d.paragraphs)
    for i in (1,200,250,251,275,300,325):
        assert f'GRV-AI-{i:03d}' in text, f'Prompt {i} missing from Word'
    assert len([p for p in d.paragraphs if p.text.startswith('MANAGER CHECK')]) == COUNT

def main():
    rows=fetch_rows()
    ordered,groups=group_sections(rows)
    index_path=INDEX/f'restaurant-ai-prompts-{COUNT}.json'
    index_path.write_text(json.dumps(rows,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    READER.parent.mkdir(parents=True,exist_ok=True)
    READER.write_text(reader_text(rows,ordered,groups),encoding='utf-8')
    docx=OUT/f'{STEM}.docx'
    pdf=OUT/f'{STEM}.pdf'
    build_docx(rows,ordered,groups,docx)
    verify_docx(docx)
    build_pdf(rows,ordered,groups,pdf)
    print('AI_WORKBOOK_PROMPTS='+str(len(rows)))
    print('AI_WORKBOOK_CATEGORIES='+str(len(ordered)))
    print('AI_WORKBOOK_ADDED=75 BASELINE_PRESERVED=250')
    print(f'AI_WORKBOOK_DOCX={docx.relative_to(ROOT)} ({docx.stat().st_size} bytes)')
    print(f'AI_WORKBOOK_PDF={pdf.relative_to(ROOT)} ({pdf.stat().st_size} bytes)')
    print(f'AI_WORKBOOK_INDEX={index_path.relative_to(ROOT)}')
    print(f'AI_WORKBOOK_READER={READER.relative_to(ROOT)} ({len(READER.read_text(encoding="utf-8"))} chars)')

if __name__=='__main__':
    main()
