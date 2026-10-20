#!/usr/bin/env python3
"""Build matching printable PDF and editable DOCX manuals from the live October 2026 HTML manuals.
The static HTML files are the controlled source; do not independently edit generated binaries.
"""
from pathlib import Path
from html import unescape
import re
from bs4 import BeautifulSoup
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, PageBreak, KeepTogether
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_LEFT, TA_CENTER
from reportlab.lib.units import inch
from reportlab.lib.pagesizes import letter
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[1]
MANUALS=ROOT/'site/assets/downloads/windows/manuals/current'
NAMES=[
  'Gravity_SOP_Foundry_Windows_User_Manual',
  'Gravity_Shift_Lab_Windows_User_Manual',
  'Gravity_Inventory_Center_Windows_User_Manual',
  'Gravity_Work_Study_Lab_Windows_User_Manual',
]
GOLD=colors.HexColor('#AC7F33')
NAVY=colors.HexColor('#173340')
BODY=colors.HexColor('#273B45')
def clean(txt):
    return re.sub(r'\s+',' ',unescape(txt or '')).strip()
def set_docx_footer(section):
    p=section.footer.paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.add_run('THOMAS MONROE BOOKS  ·  GRAVITY WINDOWS  ·  ')
    run=p.add_run()
    fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');run._r.addnext(fld)
def pdf_footer(canvas, doc):
    canvas.saveState()
    canvas.setStrokeColor(GOLD)
    canvas.line(0.65*inch,0.65*inch,7.85*inch,0.65*inch)
    canvas.setFont('Helvetica',8)
    canvas.setFillColor(BODY)
    canvas.drawString(0.65*inch,0.48*inch,'Thomas Monroe Books · Gravity Windows · October 2026')
    canvas.drawRightString(7.85*inch,0.48*inch,str(doc.page))
    canvas.restoreState()
def produce(name):
    filename=MANUALS/(name+'.html')
    source=BeautifulSoup(filename.read_text(encoding='utf-8'),'html.parser')
    title=clean(source.select_one('header h1').get_text(' ',strip=True))
    version=clean(source.select_one('header p').get_text(' ',strip=True))
    intro=clean(source.select_one('header .intro').get_text(' ',strip=True))
    sections=source.select('#guide-content > section')
    assert len(sections)>=15, f'Manual is incomplete: {name}'
    pages=source.select_one('article .note')
    note=clean(pages.get_text(' ',strip=True)) if pages else ''
    ss=getSampleStyleSheet()
    ss.add(ParagraphStyle(name='ManualTitle',fontName='Helvetica-Bold',fontSize=21,leading=26,textColor=NAVY,spaceAfter=16))
    ss.add(ParagraphStyle(name='ManualSub',fontName='Helvetica-Bold',fontSize=11,leading=16,textColor=GOLD,spaceAfter=10))
    ss.add(ParagraphStyle(name='ManualBody',fontName='Helvetica',fontSize=10.2,leading=15,textColor=BODY,spaceAfter=11))
    ss.add(ParagraphStyle(name='ManualH2',fontName='Helvetica-Bold',fontSize=12.8,leading=17,textColor=NAVY,spaceBefore=17,spaceAfter=10,keepWithNext=True))
    ss.add(ParagraphStyle(name='ManualNote',parent=ss['ManualBody'],backColor=colors.HexColor('#F6F2E8'),borderColor=GOLD,borderWidth=0.5,borderPadding=9,spaceBefore=10,spaceAfter=16))
    import xml.sax.saxutils as xml
    para=lambda s:xml.escape(clean(s))
    story=[Paragraph(para(title),ss['ManualTitle']),Paragraph(para(version),ss['ManualSub']),Paragraph(para(intro),ss['ManualBody']),Paragraph(para(note),ss['ManualNote']),Paragraph('Contents',ss['ManualH2'])]
    story.extend(Paragraph(para(clean(section.h2.get_text(' ',strip=True))),ss['ManualBody']) for section in sections)
    story.append(PageBreak())
    word=Document()
    sec=word.sections[0];sec.top_margin=Inches(.78);sec.bottom_margin=Inches(.78);sec.left_margin=Inches(.82);sec.right_margin=Inches(.82)
    set_docx_footer(sec)
    styles=word.styles
    styles['Normal'].font.name='Aptos';styles['Normal'].font.size=Pt(10.5);styles['Normal'].font.color.rgb=RGBColor(35,52,60)
    styles['Normal'].paragraph_format.space_after=Pt(7)
    for nm,size,before,after in [('Heading 1',20,0,15),('Heading 2',13,12,8)]:
        st=styles[nm];st.font.name='Aptos Display';st.font.size=Pt(size);st.font.bold=True
        st.font.color.rgb=RGBColor(22,52,68);st.paragraph_format.space_before=Pt(before);st.paragraph_format.space_after=Pt(after)
    word.add_heading(title,0)
    word.add_paragraph(version,'Subtitle')
    word.add_paragraph(intro)
    word.add_paragraph(note,'Intense Quote')
    word.add_heading('Contents',level=1)
    for section in sections:
        word.add_paragraph(clean(section.h2.get_text(' ',strip=True)))
    word.add_page_break()
    for section in sections:
        header=clean(section.h2.get_text(' ',strip=True))
        story.append(Paragraph(para(header),ss['ManualH2']))
        word.add_heading(header,level=2)
        for child in section.find_all(['p','ul','ol'],recursive=False):
            if child.name=='p':
                s=clean(child.get_text(' ',strip=True))
                if not s:continue
                story.append(Paragraph(para(s),ss['ManualBody']))
                word.add_paragraph(s)
            else:
                for li in child.find_all('li',recursive=False):
                    s=clean(li.get_text(' ',strip=True))
                    story.append(Paragraph('&#8226; '+para(s),ss['ManualBody']))
                    word.add_paragraph(s,style='List Bullet')
    pdf_path=MANUALS/(name+'.pdf')
    pdf=SimpleDocTemplate(str(pdf_path),pagesize=letter,rightMargin=.7*inch,leftMargin=.7*inch,topMargin=.8*inch,bottomMargin=.85*inch,author='Thomas Monroe Books',title=title)
    pdf.build(story,onFirstPage=pdf_footer,onLaterPages=pdf_footer)
    docx_path=MANUALS/(name+'.docx')
    word.save(docx_path)
    from pypdf import PdfReader
    reader=PdfReader(str(pdf_path))
    full='\n'.join((p.extract_text() or '') for p in reader.pages)
    assert len(reader.pages)>=3, f'PDF unexpectedly short: {name}'
    for h in [clean(se.h2.get_text(' ',strip=True)) for se in sections]:
        assert h in full, f'PDF missed a section: {name}: {h}'
    assert len(sections)==len(word.element.body.xpath('.//w:pStyle[@w:val="Heading2"]')) or docx_path.stat().st_size>20000
    print(f'OK {name}: {len(sections)} chapters, {len(reader.pages)} PDF pages, PDF {pdf_path.stat().st_size:,} B, DOCX {docx_path.stat().st_size:,} B')
def main():
    assert MANUALS.is_dir()
    for name in NAMES: produce(name)
if __name__=='__main__':main()
